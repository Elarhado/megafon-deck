#!/usr/bin/env python3
"""Build a MegaFon-branded .pptx from deck.json.  Usage: build_deck.py deck.json -o out.pptx

Schema: see deck.schema.md.  Exit codes: 0 ok, 2 validation error, 3 missing dependency.
"""
import argparse
import copy
import difflib
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _common as C  # noqa: E402

C.require("pptx", "PIL")

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402
from lxml import etree  # noqa: E402

STYLES = ("light", "green", "dark")
STYLE_INDEX = {"dark": 0, "green": 1, "light": 2}
SLIDE_KEYS = {"layout", "style", "notes", "speaker_notes", "elements"}


def slide_keys(layout):
    """Reserved slide keys (speaker notes: 'notes' or 'speaker_notes')."""
    keys = set(SLIDE_KEYS)
    if "notes" in ROLES[layout]["roles"]:
        keys.discard("notes")
    return keys


def speaker_notes(spec):
    return spec.get("speaker_notes") or spec.get("notes")
CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
ICON_COLORS = ("green", "purple", "white", "black", "berry")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

BRAND = C.load_json("brand.json")
LAYOUTS = C.load_json("layouts.json")
ROLES = C.load_json("layout_roles.json")["layouts"]
_, _, NAMED = C.palette()


class ValidationError(Exception):
    pass


# ------------------------------------------------------------------ validation
def icon_names():
    d = os.path.join(C.ASSETS, "icons", "green")
    return sorted(f[:-4] for f in os.listdir(d) if f.endswith(".png"))


def role_slots(layout, role):
    """Return (idxs list, kind) where kind is 'scalar', 'list' or 'matrix'."""
    v = ROLES[layout]["roles"][role]
    if isinstance(v, int):
        return [v], "scalar"
    if v and isinstance(v[0], list):
        return v, "matrix"
    return v, "list"


def check_icon(name, color, where, errs):
    if color not in ICON_COLORS:
        errs.append("%s: icon colour '%s' not in %s" % (where, color, ", ".join(ICON_COLORS)))
        return
    if name not in icon_names():
        close = difflib.get_close_matches(name, icon_names(), n=10, cutoff=0.0)
        errs.append("%s: unknown icon '%s'. Closest: %s" % (where, name, ", ".join(close)))


def check_chart(ch, where, errs):
    if not isinstance(ch, dict):
        errs.append("%s: chart must be an object" % where)
        return
    if ch.get("type") not in CHART_TYPES:
        errs.append("%s: chart.type must be one of %s" % (where, ", ".join(CHART_TYPES)))
    cats, ser = ch.get("categories"), ch.get("series")
    if not cats or not isinstance(cats, list):
        errs.append("%s: chart.categories missing" % where)
        return
    if not ser or not isinstance(ser, list):
        errs.append("%s: chart.series missing" % where)
        return
    for s in ser:
        vals = s.get("values") if isinstance(s, dict) else None
        if not vals or len(vals) != len(cats):
            errs.append("%s: series '%s' needs %d values" % (where, (s or {}).get("name", "?"), len(cats)))


def validate(deck):
    errs = []
    if not isinstance(deck.get("slides"), list) or not deck["slides"]:
        raise ValidationError("deck.slides must be a non-empty list")
    if deck.get("style", "light") not in STYLES:
        errs.append("deck.style must be one of %s" % ", ".join(STYLES))
    for n, s in enumerate(deck["slides"], 1):
        w = "slide %d" % n
        lay = s.get("layout")
        if lay not in ROLES:
            errs.append("%s: unknown layout '%s'. Layouts: %s" % (w, lay, ", ".join(ROLES)))
            continue
        if s.get("style", deck.get("style", "light")) not in STYLES:
            errs.append("%s: style must be one of %s" % (w, ", ".join(STYLES)))
        allowed = list(ROLES[lay]["roles"])
        for key, val in s.items():
            if key in slide_keys(lay):
                continue
            if key not in ROLES[lay]["roles"]:
                errs.append("%s (%s): unknown role '%s'. Allowed roles: %s" % (w, lay, key, ", ".join(allowed)))
                continue
            if key == "chart":
                check_chart(val, "%s.chart" % w, errs)
                continue
            if key == "picture":
                if not isinstance(val, dict) or not ("icon" in val or "image" in val):
                    errs.append('%s.picture must be {"icon","color"} or {"image"}' % w)
                elif "icon" in val:
                    check_icon(val["icon"], val.get("color", "green"), w + ".picture", errs)
                elif not os.path.exists(val["image"]):
                    errs.append("%s.picture: image file not found: %s" % (w, val["image"]))
                continue
            idxs, kind = role_slots(lay, key)
            if kind == "scalar":
                if isinstance(val, list) and not all(isinstance(x, str) for x in val):
                    errs.append("%s.%s: expected text, list of bullets or {text,bullets}" % (w, key))
            elif not isinstance(val, list):
                errs.append("%s.%s: expected a list (max %d items)" % (w, key, len(idxs)))
            elif len(val) > len(idxs):
                errs.append("%s.%s: %d items but layout '%s' has only %d slots" % (w, key, len(val), lay, len(idxs)))
            elif kind == "matrix":
                for r in val:
                    if not isinstance(r, list) or len(r) > len(idxs[0]):
                        errs.append("%s.%s: each row must be a list of up to %d cells" % (w, key, len(idxs[0])))
                        break
        if "chart" in ROLES[lay]["roles"] and "chart" not in s:
            errs.append("%s (%s): missing chart data. Add \"chart\": {type, categories, series}" % (w, lay))
        for i, el in enumerate(s.get("elements", []), 1):
            ew = "%s.elements[%d]" % (w, i)
            kind = el.get("kind")
            if kind not in ("chart", "table", "text", "icon", "image"):
                errs.append("%s: kind must be chart|table|text|icon|image" % ew)
            elif kind == "chart":
                check_chart(el.get("chart"), ew + ".chart", errs)
            elif kind == "table" and not el.get("rows"):
                errs.append("%s: table needs rows" % ew)
            elif kind == "icon":
                check_icon(el.get("name", ""), el.get("color", "green"), ew, errs)
            elif kind == "image" and not os.path.exists(el.get("path", "")):
                errs.append("%s: image file not found: %s" % (ew, el.get("path")))
            if kind in ("chart", "table", "text", "image") and not all(k in el for k in ("x", "y", "w", "h")):
                errs.append("%s: x, y, w, h required" % ew)
            if kind == "icon" and not all(k in el for k in ("x", "y", "size")):
                errs.append("%s: x, y, size required" % ew)
    if errs:
        raise ValidationError("\n".join(errs))


# --------------------------------------------------------------------- helpers
def ts(x):
    return C.typeset(x) if isinstance(x, str) else x


def color_hex(c, style):
    if c is None:
        return BRAND["text_colors"][style]
    c = str(c).lstrip("#")
    return NAMED.get(c.lower(), c).upper()


def default_icon_color(style):
    return "white" if style == "green" else "green"


def alpha_color(parent_el, hexval, alpha_pct):
    """Append <a:solidFill><a:srgbClr val=hex><a:alpha/></...> (brand colour with transparency)."""
    sf = etree.SubElement(parent_el, qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"), val=hexval)
    if alpha_pct < 100:
        etree.SubElement(c, qn("a:alpha"), val=str(int(alpha_pct * 1000)))
    return sf


def line_xml(hexval, alpha, width_pt, tag="a:ln"):
    ln = etree.Element(qn(tag), w=str(int(width_pt * 12700)))
    alpha_color(ln, hexval, alpha)
    return ln


def grid_colour(style):
    return ("000000", 15) if style == "light" else ("FFFFFF", 25)


# ------------------------------------------------------------------ text filling
def paragraphs_of(value):
    """Normalise a role value to a list of (text, level) paragraphs."""
    if isinstance(value, str):
        return [(ts(t), 0) for t in value.split("\n")]
    if isinstance(value, dict):
        out = [(ts(value.get("text", "")), 0)]
        out += [(ts(b), 1) for b in value.get("bullets", [])]
        return out
    out = []
    for item in value:
        out += paragraphs_of(item)
    return out


def has_level1_style(ph):
    ls = ph._element.find(".//" + qn("a:lstStyle"))
    return ls is not None and ls.find(qn("a:lvl2pPr")) is not None


def fill_text(ph, value):
    """Fill a placeholder; returns False (and leaves it untouched) when the value is empty."""
    paras = paragraphs_of(value)
    if not any(t.strip() for t, _ in paras):
        return False
    use_level = has_level1_style(ph)
    tf = ph.text_frame
    for i, (text, lvl) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if lvl and use_level:
            p.level = 1
        p.add_run().text = text
    return True


def ph_by_idx(slide):
    return {sh.placeholder_format.idx: sh for sh in slide.placeholders}


def remove_shape(sh):
    sh._element.getparent().remove(sh._element)


# ------------------------------------------------------------------------ charts
def series_colours(spec, style, layout_key=None):
    """Brand colour order for a chart: the layout's legend dots if it has them, green for a lone series."""
    by_layout = BRAND.get("chart_series_by_layout", {})
    if layout_key in by_layout and not layout_key.startswith("_"):
        return by_layout[layout_key]
    if spec["type"] not in ("pie", "doughnut") and len(spec.get("series", [])) == 1:
        return [BRAND.get("chart_single_series", {}).get(style, "00B956")]
    return BRAND["chart_series"][style]


def style_chart(chart, spec, style, pt, layout_key=None):
    from pptx.enum.chart import XL_TICK_LABEL_POSITION  # noqa: F401
    ctype = spec["type"]
    text = RGBColor.from_string(BRAND["text_colors"][style])
    series_cols = spec.get("colors") or series_colours(spec, style, layout_key)
    chart.font.name = "Arial"
    chart.font.size = Pt(pt)
    chart.font.color.rgb = text
    chart.has_title = False
    chart.has_legend = bool(spec.get("legend", False))
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(pt)
        chart.legend.font.name = "Arial"
        chart.legend.font.color.rgb = text
    plot = chart.plots[0]
    pie = ctype in ("pie", "doughnut")
    if pie:
        plot.vary_by_categories = True
        for i, pnt in enumerate(plot.series[0].points):
            pnt.format.fill.solid()
            pnt.format.fill.fore_color.rgb = RGBColor.from_string(series_cols[i % len(series_cols)])
            pnt.format.line.color.rgb = RGBColor.from_string("FFFFFF" if style != "dark" else "000000")
    else:
        plot.vary_by_categories = False
        for i, s in enumerate(plot.series):
            col = RGBColor.from_string(series_cols[i % len(series_cols)])
            if ctype == "line":
                s.format.line.color.rgb = col
                s.format.line.width = Pt(2.25)
                s.smooth = False
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = col
                s.marker.format.line.color.rgb = col
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = col
                s.format.line.fill.background()
    if ctype in ("bar", "column", "stacked_column"):
        plot.gap_width = 60
        if ctype == "stacked_column":
            plot.overlap = 100
    if spec.get("labels", True):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = spec.get("number_format", "General")
        dl.number_format_is_linked = False
        dl.show_value = True
        dl.font.size = Pt(pt)
        dl.font.name = "Arial"
        dl.font.color.rgb = text
        if pie:
            sdl = plot.series[0].data_labels  # series-level labels override plot-level ones
            sdl.show_value = True
            sdl.number_format = spec.get("number_format", "General")
            sdl.number_format_is_linked = False
            sdl.font.size = Pt(pt)
            sdl.font.name = "Arial"
            sdl.font.color.rgb = text
            if ctype == "pie":
                sdl.position = XL_LABEL_POSITION.OUTSIDE_END
            for i, pnt in enumerate(plot.series[0].points):
                hexv = series_cols[i % len(series_cols)]
                r_, g_, b_ = (int(hexv[k:k + 2], 16) for k in (0, 2, 4))
                if 0.299 * r_ + 0.587 * g_ + 0.114 * b_ < 110:  # dark slice -> white label
                    pdl = pnt.data_label
                    pdl.font.size = Pt(pt)
                    pdl.font.name = "Arial"
                    pdl.font.color.rgb = RGBColor.from_string("FFFFFF")
                    el = pdl._dLbl
                    nf = etree.Element(qn("c:numFmt"), formatCode=spec.get("number_format", "General"),
                                       sourceLinked="0")
                    anchor = el.find(qn("c:spPr"))
                    if anchor is None:
                        anchor = el.find(qn("c:txPr"))
                    anchor.addprevious(nf)
        if ctype in ("bar", "column"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        elif ctype == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
        elif ctype == "pie":
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        elif ctype == "stacked_column":
            dl.position = XL_LABEL_POSITION.CENTER  # doughnut: no position allowed by the format
    if not pie:
        gcol, galpha = grid_colour(style)
        va, ca = chart.value_axis, chart.category_axis
        va.has_major_gridlines = True
        gl_el = va._element.find(qn("c:majorGridlines"))
        for old in gl_el.findall(qn("c:spPr")):
            gl_el.remove(old)
        sp = etree.SubElement(gl_el, qn("c:spPr"))
        sp.append(line_xml(gcol, galpha, 0.5))
        va.format.line.fill.background()
        va.has_title = False
        ca.has_title = False
        va.tick_labels.font.size = Pt(pt)
        ca.tick_labels.font.size = Pt(pt)
        if spec.get("labels", True):
            va.visible = False  # values are printed on the data labels
        ca.format.line.color.rgb = RGBColor.from_string(BRAND["text_colors"][style])
        ca.format.line.width = Pt(0.5)
        if ctype == "bar":
            ca.reverse_order = True
    # no chart border / fill
    cs = chart._chartSpace
    for old in cs.findall(qn("c:spPr")):
        cs.remove(old)
    sp = etree.Element(qn("c:spPr"))
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    cs.find(qn("c:chart")).addnext(sp)
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    for old in pa.findall(qn("c:spPr")):
        pa.remove(old)
    psp = etree.SubElement(pa, qn("c:spPr"))
    etree.SubElement(psp, qn("a:noFill"))
    ln = etree.SubElement(psp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def chart_data(spec):
    cd = CategoryChartData()
    cd.categories = [ts(str(c)) for c in spec["categories"]]
    for s in spec["series"]:
        cd.add_series(ts(s.get("name", "")), s["values"])
    return cd


def add_chart_to_placeholder(ph, spec, style, pt, layout_key=None):
    gf = ph.insert_chart(CHART_TYPES[spec["type"]], chart_data(spec))
    style_chart(gf.chart, spec, style, pt, layout_key)
    return gf


# ------------------------------------------------------------------------ icons
def icon_path(name, color):
    return os.path.join(C.ASSETS, "icons", color, name + ".png")


def insert_picture_role(slide, ph, spec, style):
    if "icon" in spec:
        path = icon_path(spec["icon"], spec.get("color", default_icon_color(style)))
        # icons are square: crop-to-fill would clip them in non-square placeholders -> fit instead
        w, h = ph.width, ph.height
        if abs(w / h - 1) < 0.05:
            ph.insert_picture(path)
        else:
            side = min(w, h)
            slide.shapes.add_picture(path, ph.left + (w - side) // 2, ph.top + (h - side) // 2, side, side)
            remove_shape(ph)
    else:
        ph.insert_picture(spec["image"])


# ------------------------------------------------------------------------ elements
def add_text_element(slide, el, style):
    tb = slide.shapes.add_textbox(Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(ts(el["text"]).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name = "Arial"
        r.font.size = Pt(el.get("pt", 14))
        r.font.bold = bool(el.get("bold", False))
        r.font.color.rgb = RGBColor.from_string(color_hex(el.get("color"), style))


def set_cell_border(cell, side, hexval, alpha, width_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"b": "a:lnB", "t": "a:lnT", "l": "a:lnL", "r": "a:lnR"}[side]
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    if hexval is None:
        ln = etree.Element(qn(tag), w="0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln = line_xml(hexval, alpha, width_pt, tag)
    # borders must precede fill elements inside tcPr
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    pos = 0
    for child in tcPr:
        if child.tag in [qn(t) for t in order[:order.index(tag)]]:
            pos = list(tcPr).index(child) + 1
    tcPr.insert(pos, ln)


def add_table_element(slide, el, style):
    rows = el["rows"]
    nrows, ncols = len(rows), max(len(r) for r in rows)
    gf = slide.shapes.add_table(nrows, ncols, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    header = el.get("header", True)
    tbl.first_row = bool(header)
    tbl.horz_banding = False
    pt = el.get("pt", 10)
    head_fill = "731982" if style == "green" else "00B956"
    text_hex = BRAND["text_colors"][style]
    gcol, galpha = grid_colour(style)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Emu(int(Inches(el["h"]) / nrows))
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            txt = ts(str(row[ci])) if ci < len(row) else ""
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            tf = cell.text_frame
            tf.word_wrap = True
            r = tf.paragraphs[0].add_run()
            r.text = txt
            r.font.name = "Arial"
            r.font.size = Pt(pt)
            is_head = header and ri == 0
            r.font.bold = bool(is_head)
            if is_head:
                r.font.color.rgb = RGBColor.from_string("FFFFFF")
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(head_fill)
            else:
                r.font.color.rgb = RGBColor.from_string(text_hex)
                cell.fill.background()
            for side in ("l", "r", "t"):
                set_cell_border(cell, side, None, 0, 0)
            set_cell_border(cell, "b", gcol, galpha + 15, 0.5)
            from pptx.enum.text import MSO_ANCHOR
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_element(slide, el, style):
    kind = el["kind"]
    if kind == "chart":
        gf = slide.shapes.add_chart(CHART_TYPES[el["chart"]["type"]], Inches(el["x"]), Inches(el["y"]),
                                    Inches(el["w"]), Inches(el["h"]), chart_data(el["chart"]))
        style_chart(gf.chart, el["chart"], style, el["chart"].get("pt", 10))
    elif kind == "table":
        add_table_element(slide, el, style)
    elif kind == "text":
        add_text_element(slide, el, style)
    elif kind == "icon":
        s = Inches(el["size"])
        slide.shapes.add_picture(icon_path(el["name"], el.get("color", default_icon_color(style))),
                                 Inches(el["x"]), Inches(el["y"]), s, s)
    elif kind == "image":
        slide.shapes.add_picture(el["path"], Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))


# ----------------------------------------------------------------- auto placeholders
def copy_auto_placeholders(slide, layout, header_label):
    """Copy date and slide-number <p:sp> from the layout into the slide (python-pptx skips them)."""
    tree = slide.shapes._spTree
    next_id = max([int(i) for i in tree.xpath("//p:cNvPr/@id")] + [1]) + 1
    n_auto = 0
    for sp in layout.placeholders:
        t = sp.placeholder_format.type
        name = str(t)
        if "DATE" in name or "SLIDE_NUMBER" in name:
            el = copy.deepcopy(sp._element)
            el.find(".//" + qn("p:cNvPr")).set("id", str(next_id))
            next_id += 1
            # slide-level shapes must not carry layout geometry overrides unless present: keep as in layout
            if "DATE" in name:
                set_header_label(el, header_label)
            tree.append(el)
            n_auto += 1
    return n_auto


def set_header_label(sp_el, label):
    """Fill the date placeholder: first run = deck name, last run = date (template layout: name <br> date)."""
    title, _, date = label.partition("\x0b")
    p = sp_el.find(".//" + qn("a:p"))
    runs = p.findall(qn("a:r"))
    br = p.find(qn("a:br"))
    if br is not None and len(runs) >= 2:
        runs[0].find(qn("a:t")).text = title
        for mid in runs[1:-1]:
            mid.find(qn("a:t")).text = ""
        runs[-1].find(qn("a:t")).text = date
        return
    for ch in list(p):
        if ch.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p.remove(ch)
    end = p.find(qn("a:endParaRPr"))
    r = etree.Element(qn("a:r"))
    etree.SubElement(r, qn("a:rPr"), lang="ru-RU", sz="800", b="1")
    etree.SubElement(r, qn("a:t")).text = title + " " + date
    if end is not None:
        end.addprevious(r)
    else:
        p.append(r)


# --------------------------------------------------------------------------- build
NOTES_BODY = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="%d" name="Notes Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" sz="quarter" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ru-RU"/></a:p></p:txBody></p:sp>')


def ensure_notes_body(prs):
    """The template's notes master has no body placeholder; add one so speaker notes work."""
    nm = prs.notes_master
    if any(sh.is_placeholder and sh.placeholder_format.idx == 1 for sh in nm.shapes):
        return
    tree = nm.shapes._spTree
    next_id = max([int(i) for i in tree.xpath("//p:cNvPr/@id")] + [1]) + 1
    tree.append(etree.fromstring(NOTES_BODY % next_id))


def build(deck, out_path):
    prs = Presentation(os.path.join(C.ASSETS, "template.pptx"))
    ensure_notes_body(prs)
    default_style = deck.get("style", "light")
    date = deck.get("date", "")
    label = ts("%s\x0b%s" % (deck.get("title", ""), date)).strip("\x0b")
    summary = []
    for n, spec in enumerate(deck["slides"], 1):
        style = spec.get("style", default_style)
        lay_key = spec["layout"]
        L = LAYOUTS["styles"][style][lay_key]
        layout = prs.slide_masters[L["master"]].slide_layouts[L["index"]]
        slide = prs.slides.add_slide(layout)
        phs = ph_by_idx(slide)
        used = set()
        roles = ROLES[lay_key]["roles"]
        for role, value in spec.items():
            if role in slide_keys(lay_key):
                continue
            idxs, kind = role_slots(lay_key, role)
            if role == "chart":
                ph = phs[idxs[0]]
                add_chart_to_placeholder(ph, value, style, next((s["pt"] for s in L["slots"] if s["idx"] == idxs[0]), 10) or 10,
                                         spec["layout"])
                used.add(idxs[0])
            elif role == "picture":
                insert_picture_role(slide, phs[idxs[0]], value, style)
                used.add(idxs[0])
            elif kind == "scalar":
                if fill_text(phs[idxs[0]], value):
                    used.add(idxs[0])
            elif kind == "list":
                for idx, item in zip(idxs, value):
                    if fill_text(phs[idx], item):
                        used.add(idx)
            else:  # matrix
                for row_idxs, row in zip(idxs, value):
                    for idx, item in zip(row_idxs, row):
                        if fill_text(phs[idx], item):
                            used.add(idx)
        if lay_key == "cover" and "label" not in spec and date:
            if fill_text(phs[11], date):
                used.add(11)
        # drop unused placeholders (empty text / unfilled pictures)
        for idx, sh in phs.items():
            if idx not in used:
                remove_shape(sh)
        n_auto = copy_auto_placeholders(slide, layout, label)
        for el in spec.get("elements", []):
            add_element(slide, el, style)
        if speaker_notes(spec):
            slide.notes_slide.notes_text_frame.text = ts(speaker_notes(spec))
        title_txt = ""
        if "title" in spec:
            title_txt = spec["title"] if isinstance(spec["title"], str) else str(spec["title"])
        elif "quote" in spec:
            title_txt = str(spec["quote"])[:40]
        summary.append("%02d %-24s [%s] %-40.40s placeholders=%d auto=%d elements=%d"
                       % (n, lay_key, style, title_txt.replace("\n", " "), len(used), n_auto,
                          len(spec.get("elements", []))))
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    return summary


def main(argv=None):
    if argv is None:
        argv = sys.argv[1:]
    if "--check" in argv:  # dependencies were imported at module load (C.require exits 3 otherwise)
        import pptx, PIL
        print("ok: python %s, python-pptx %s, Pillow %s" % (sys.version.split()[0], pptx.__version__, PIL.__version__))
        return 0
    ap = argparse.ArgumentParser(description="Build MegaFon-branded pptx from deck.json")
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args(argv)
    try:
        with open(a.deck, encoding="utf-8") as f:
            deck = json.load(f)
        validate(deck)
    except (ValidationError, json.JSONDecodeError, OSError) as e:
        print("Validation error:\n" + str(e), file=sys.stderr)
        return 2
    for line in build(deck, a.out):
        print(line)
    print("saved %s" % a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
