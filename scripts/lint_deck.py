#!/usr/bin/env python3
"""Lint a MegaFon deck.  Usage: lint_deck.py deck.pptx [--json]   (exit 1 when errors are found)"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _common as C  # noqa: E402

C.require("pptx", "PIL")

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from PIL import ImageFont  # noqa: E402

EMU_IN = 914400
LAYOUTS = C.load_json("layouts.json")
ROLES = C.load_json("layout_roles.json")["layouts"]
ALLOWED, GRADIENT_ONLY, _ = C.palette()
PROMPT_EXACT = {"текст", "заголовок", "подзаголовок", "образец заголовка", "название презентации"}
PROMPT_START = ("click to add", "образец ", "on the other hand", "lorem ipsum")
STYLE_BY_MASTER = {v: k for k, v in {"dark": 0, "green": 1, "light": 2}.items()}


class Issues:
    def __init__(self):
        self.items = []

    def add(self, slide, shape, level, msg):
        self.items.append({"slide": slide, "shape": shape, "level": level, "message": msg})


# ------------------------------------------------------------- font measurement
_font_cache = {}


def measurer(bold):
    path = C.find_font(bold)
    if path is None:
        return lambda text, pt: len(text) * pt * (0.56 if bold else 0.5)
    def width(text, pt):
        key = (bold,)
        f = _font_cache.get(key)
        if f is None:
            f = _font_cache[key] = ImageFont.truetype(path, 100)
        return f.getlength(text) * pt / 100.0
    return width


MEASURE = {False: measurer(False), True: measurer(True)}


def para_text(p):
    out = []
    for ch in p:
        if ch.tag == qn("a:r"):
            out.append(ch.findtext(qn("a:t")) or "")
        elif ch.tag == qn("a:br"):
            out.append("\n")
        elif ch.tag == qn("a:fld"):
            out.append(ch.findtext(qn("a:t")) or "0")
    return "".join(out)


def shape_text(sh):
    if not sh.has_text_frame:
        return ""
    return "\n".join(para_text(p._p) for p in sh.text_frame.paragraphs)


# ----------------------------------------------------- inheritance for text size
def layout_shape(sh, layout):
    if not sh.is_placeholder:
        return None
    for l in layout.placeholders:
        if l.placeholder_format.idx == sh.placeholder_format.idx:
            return l
    return None


def style_chain(sh, layout, prs_master):
    """XML elements (lvlNpPr parents) that may define defaults, most specific first."""
    chain = []
    ls = sh._element.find(".//" + qn("a:lstStyle"))
    if ls is not None:
        chain.append(ls)
    lsh = layout_shape(sh, layout) if sh.is_placeholder else None
    if lsh is not None:
        l2 = lsh._element.find(".//" + qn("a:lstStyle"))
        if l2 is not None:
            chain.append(l2)
        kind = "title" if lsh.placeholder_format.type is not None and "TITLE" in str(lsh.placeholder_format.type) else "body"
    else:
        kind = "other"
    tx = prs_master._element.find(".//" + qn("p:txStyles"))
    if tx is not None:
        node = tx.find(qn({"title": "p:titleStyle", "body": "p:bodyStyle"}.get(kind, "p:otherStyle")))
        if node is not None:
            chain.append(node)
    return chain


def lvl_attr(chain, lvl, getter):
    for ls in chain:
        pp = ls.find(qn("a:lvl%dpPr" % (lvl + 1)))
        if pp is not None:
            v = getter(pp)
            if v is not None:
                return v
    return None


def eff_font(sh, layout, master, slot_pt, p, r):
    """Effective (size pt, bold, caps, line spacing, space-before pt, marL in, black) for a run."""
    chain = style_chain(sh, layout, master)
    lvl = int(p.get("lvl", 0))
    sz = None
    if r is not None and r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).get("sz"):
        sz = int(r.find(qn("a:rPr")).get("sz")) / 100
    if sz is None:
        sz = slot_pt
    if sz is None:
        sz = lvl_attr(chain, lvl, lambda pp: (int(pp.find(qn("a:defRPr")).get("sz")) / 100
                                              if pp.find(qn("a:defRPr")) is not None
                                              and pp.find(qn("a:defRPr")).get("sz") else None)) or 18
    def rattr(name):
        if r is not None and r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).get(name) is not None:
            return r.find(qn("a:rPr")).get(name)
        return lvl_attr(chain, lvl, lambda pp: (pp.find(qn("a:defRPr")).get(name)
                                                if pp.find(qn("a:defRPr")) is not None else None))
    bold = rattr("b") in ("1", "true")
    caps = rattr("cap") == "all"
    latin = None
    if r is not None and r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).find(qn("a:latin")) is not None:
        latin = r.find(qn("a:rPr")).find(qn("a:latin")).get("typeface")
    if latin is None:
        latin = lvl_attr(chain, lvl, lambda pp: (pp.find(qn("a:defRPr")).find(qn("a:latin")).get("typeface")
                                                 if pp.find(qn("a:defRPr")) is not None
                                                 and pp.find(qn("a:defRPr")).find(qn("a:latin")) is not None else None))
    black = bool(latin and "black" in latin.lower())
    ppr = p.find(qn("a:pPr"))

    def ppr_val(tag, sub):
        if ppr is not None and ppr.find(qn(tag)) is not None and ppr.find(qn(tag)).find(qn(sub)) is not None:
            return ppr.find(qn(tag)).find(qn(sub)).get("val")
        return lvl_attr(chain, lvl, lambda pp: (pp.find(qn(tag)).find(qn(sub)).get("val")
                                                if pp.find(qn(tag)) is not None
                                                and pp.find(qn(tag)).find(qn(sub)) is not None else None))
    ln = ppr_val("a:lnSpc", "a:spcPct")
    sb = ppr_val("a:spcBef", "a:spcPts")
    marl = None
    if ppr is not None and ppr.get("marL"):
        marl = int(ppr.get("marL"))
    if marl is None:
        marl = lvl_attr(chain, lvl, lambda pp: int(pp.get("marL")) if pp.get("marL") else None) or 0
    return (sz, bold, caps, int(ln) / 100000 if ln else 1.0, int(sb) / 100 if sb else 0.0, marl / EMU_IN, black)


def body_insets(sh, layout):
    """(l, t, r, b) insets in inches, from slide or layout bodyPr, default 0.1/0.05."""
    vals = {"lIns": 91440, "tIns": 45720, "rIns": 91440, "bIns": 45720}
    srcs = []
    lsh = layout_shape(sh, layout)
    if lsh is not None:
        srcs.append(lsh._element.find(".//" + qn("a:bodyPr")))
    srcs.append(sh._element.find(".//" + qn("a:bodyPr")))
    autofit = None
    for bp in srcs:
        if bp is None:
            continue
        for k in vals:
            if bp.get(k) is not None:
                vals[k] = int(bp.get(k))
        for ch in bp:
            if ch.tag in (qn("a:normAutofit"), qn("a:spAutoFit"), qn("a:noAutofit")):
                autofit = ch
    return vals["lIns"] / EMU_IN, vals["tIns"] / EMU_IN, vals["rIns"] / EMU_IN, vals["bIns"] / EMU_IN, autofit


def wrap_lines(text, width_pt, size, bold, caps, factor):
    m = MEASURE[bold]
    total = 0
    for hard in text.split("\n"):
        if caps:
            hard = hard.upper()
        words = hard.split(" ")
        cur = ""
        lines = 1
        for w in words:
            trial = (cur + " " + w) if cur else w
            if m(trial, size) * factor <= width_pt:
                cur = trial
            else:
                if cur:
                    lines += 1
                cur = w
                # very long word: breaks across lines
                ww = m(w, size) * factor
                while ww > width_pt:
                    lines += 1
                    ww -= width_pt
        total += lines
    return total


def needed_height(sh, layout, master, slot_pt):
    """Estimated inch height required by the shape's text, and inner width."""
    li, ti, ri, bi, autofit = body_insets(sh, layout)
    width_in = sh.width / EMU_IN - li - ri
    height = 0.0  # text height only: layout insets are not counted (single big lines are tolerated)
    first = True
    for p in sh.text_frame.paragraphs:
        pe = p._p
        text = para_text(pe)
        runs = pe.findall(qn("a:r"))
        if not text.strip() and not runs:
            # empty paragraph: still takes a line
            sz = eff_font(sh, layout, master, slot_pt, pe, None)[0]
            height += sz * 1.2 / 72
            continue
        sz, bold, caps, ln, sb, marl, black = eff_font(sh, layout, master, slot_pt, pe, runs[0] if runs else None)
        factor = 1.12 if black else 1.0
        w_pt = max((width_in - marl) * 72, 1)
        n = wrap_lines(text, w_pt, sz, bold or black, caps, factor)
        height += n * sz * 1.17 * ln / 72
        if not first:
            height += sb / 72
        first = False
    return height, width_in, autofit


# --------------------------------------------------------------------- the lint
def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type is not None and "GROUP" in str(sh.shape_type):
            yield from iter_shapes(sh.shapes)


def xml_checks(el, name, n, issues):
    for lat in el.iter(qn("a:latin")):
        tf = lat.get("typeface", "")
        if tf and not tf.startswith("+") and tf != "Arial":
            issues.add(n, name, "ERROR", "font '%s' (only Arial allowed)" % tf)
    for clr in el.iter(qn("a:srgbClr")):
        v = clr.get("val", "").upper()
        if v not in ALLOWED:
            issues.add(n, name, "ERROR", "colour #%s is not in the brand palette" % v)
        elif v in GRADIENT_ONLY:
            p, in_grad = clr.getparent(), False
            while p is not None:
                if p.tag == qn("a:gradFill"):
                    in_grad = True
                    break
                p = p.getparent()
            if not in_grad:
                issues.add(n, name, "ERROR", "colour #%s is gradient-only (brandbook pp.24-25)" % v)


def layout_key(slide, prs):
    layout = slide.slide_layout
    for mi, m in enumerate(prs.slide_masters):
        for li, l in enumerate(m.slide_layouts):
            if l.part is layout.part:
                style = STYLE_BY_MASTER[mi]
                for key, d in LAYOUTS["styles"][style].items():
                    if d["master"] == mi and d["index"] == li:
                        return key, mi, style
    return None, None, None


def lint(path):
    try:
        prs = Presentation(path)
    except Exception as e:  # missing file, not a pptx
        print("Cannot open %s: %s" % (path, e))
        sys.exit(2)
    issues = Issues()
    for n, slide in enumerate(prs.slides, 1):
        key, mi, style = layout_key(slide, prs)
        layout = slide.slide_layout
        master = layout.slide_master
        slots = {}
        if key:
            slots = {s["idx"]: s for s in LAYOUTS["styles"][style][key]["slots"]}
        has_title_shape = False
        for sh in iter_shapes(slide.shapes):
            name = sh.name + (" [idx %d]" % sh.placeholder_format.idx if sh.is_placeholder else "")
            xml_checks(sh._element, name, n, issues)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                xml_checks(sh.chart._chartSpace, name + " (chart)", n, issues)
            idx = sh.placeholder_format.idx if sh.is_placeholder else None
            ptype = str(sh.placeholder_format.type) if sh.is_placeholder else ""
            text = shape_text(sh) if sh.has_text_frame else ""
            if "TITLE" in ptype and text.strip():
                has_title_shape = True
            # empty placeholders
            if sh.is_placeholder and "SLIDE_NUMBER" not in ptype:
                is_pic = "PICTURE" in ptype or sh.shape_type is not None and "PICTURE" in str(sh.shape_type)
                if sh._element.tag == qn("p:sp") and not text.strip():
                    issues.add(n, name, "ERROR", "empty placeholder (idx %s)" % idx)
                    continue
            if not sh.has_text_frame:
                continue
            flat = text.replace("\n", " ").strip()
            low = flat.lower()
            prompt = slots.get(idx, {}).get("prompt", "").replace("\x0b", " ").strip() if idx is not None else ""
            if flat and (low in PROMPT_EXACT or low.startswith(PROMPT_START) or (len(prompt) > 12 and flat == prompt)):
                issues.add(n, name, "ERROR", "leftover prompt text: '%s'" % flat[:40])
            if not flat:
                continue
            for msg in C.typeset_issues(text):
                issues.add(n, name, "WARN", "typesetting: " + msg)
            # bullets
            if sh.is_placeholder and "TITLE" not in ptype and "DATE" not in ptype:
                npar = sum(1 for p in sh.text_frame.paragraphs if para_text(p._p).strip())
                if npar > 6:
                    issues.add(n, name, "WARN", "%d bullets in one body (max 6)" % npar)
            # overflow
            if "DATE" in ptype or "SLIDE_NUMBER" in ptype:
                continue
            try:
                need, _, autofit = needed_height(sh, layout, master, slots.get(idx, {}).get("pt"))
            except Exception as e:  # estimation must never break the lint
                issues.add(n, name, "WARN", "overflow estimate failed: %s" % e)
                continue
            if autofit is not None and autofit.tag == qn("a:spAutoFit"):
                continue
            have = sh.height / EMU_IN
            single = len([p for p in sh.text_frame.paragraphs if para_text(p._p).strip()]) == 1 and \
                need < 1.6 * (slots.get(idx, {}).get("pt") or 18) / 72 * 1.0
            if need > have * (1.35 if single else 1.05):
                issues.add(n, name, "ERROR", "text overflow: needs ~%.2f in, box is %.2f in" % (need, have))
        # tables: font/colour already scanned; check text of cells for typesetting
        for sh in iter_shapes(slide.shapes):
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for msg in C.typeset_issues(cell.text_frame.text):
                            issues.add(n, sh.name, "WARN", "typesetting in table: " + msg)
        # missing title
        if n > 1 and key and "title" in ROLES[key]["roles"] and not ROLES[key].get("no_common_title") \
                and not has_title_shape:
            issues.add(n, "-", "ERROR", "slide has no title (layout '%s' has one)" % key)
    return issues.items


def main(argv=None):
    ap = argparse.ArgumentParser(description="Lint a MegaFon pptx")
    ap.add_argument("pptx")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    items = lint(a.pptx)
    errs = sum(1 for i in items if i["level"] == "ERROR")
    warns = len(items) - errs
    if a.json:
        print(json.dumps({"errors": errs, "warnings": warns, "issues": items}, ensure_ascii=False, indent=1))
    else:
        for i in items:
            print("slide %d | %s | %s | %s" % (i["slide"], i["shape"], i["level"], i["message"]))
        print("summary: %d error(s), %d warning(s)" % (errs, warns))
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
