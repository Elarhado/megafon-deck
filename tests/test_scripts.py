"""Unit tests for build_deck / lint_deck.  Run: python3 -m unittest discover -s plugin/tests"""
import json
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
sys.path.insert(0, SCRIPTS)

import _common as C  # noqa: E402
import build_deck as B  # noqa: E402
import lint_deck as LD  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

LAYOUTS = C.load_json("layouts.json")
ROLES = C.load_json("layout_roles.json")["layouts"]
NN, NB = " ", " "


class TypesetTests(unittest.TestCase):
    def test_quotes(self):
        self.assertEqual(C.typeset('Тариф "Максимум" для всех'), "Тариф «Максимум» для всех")
        self.assertEqual(C.typeset('Он сказал: "Это \"тест\" сети"'), "Он сказал: «Это „тест“ сети»")

    def test_dash_and_ranges(self):
        self.assertEqual(C.typeset("5G - это будущее"), "5G" + NB + "— это будущее")
        self.assertEqual(C.typeset("с 5-6 октября"), "с 5–6 октября")
        self.assertEqual(C.typeset("Wi-Fi и 4G-сети"), "Wi-Fi и 4G-сети")

    def test_numbers(self):
        self.assertEqual(C.typeset("12500 ₽"), "12" + NN + "500" + NN + "₽")
        self.assertEqual(C.typeset("1500 руб."), "1" + NN + "500 руб.")
        self.assertEqual(C.typeset("В 2027 году"), "В 2027 году")
        self.assertEqual(C.typeset("1000000 абонентов"), "1" + NN + "000" + NN + "000 абонентов")
        self.assertEqual(C.typeset("12 500 ₽"), "12" + NN + "500" + NN + "₽")

    def test_untouched(self):
        for t in ("14.10.2026", "3,14159 и 150000,5", "https://megafon.ru/a-b/12345?q=5-6", "2026-10-14"):
            self.assertEqual(C.typeset(t), t)

    def test_brand(self):
        self.assertEqual(C.typeset("Мегафона, МЕГАФОН, мегафону, Megafon"), "МегаФона, МегаФон, МегаФону, MegaFon")
        self.assertEqual(C.typeset("сайт www.megafon.ru"), "сайт www.megafon.ru")

    def test_spaces_and_idempotent(self):
        self.assertEqual(C.typeset("а  б   в"), "а б в")
        s = C.typeset('"Мегафон" - 5-6 тыс. 12500 ₽')
        self.assertEqual(C.typeset(s), s)


PHRASES = ["Рост трафика", "Рост трафика 5G в городах", "Базовые станции 5G в городах-миллионниках",
           "ARPU растёт за счёт тарифов на 5G и дополнительных услуг в плотных городских зонах, а базовые станции "
           "строятся там, где нагрузка максимальна"]


def fit_text(slot):
    """Longest realistic phrase that fits the slot (conservative estimate)."""
    cap_line = max(int((slot["w"] - 0.3) * 72 / ((slot["pt"] or 14) * 0.7)), 3)
    lines = max(int(slot["h"] * 72 / ((slot["pt"] or 14) * 1.25)), 1)
    cap = cap_line * lines
    best = "5G"
    for p in PHRASES:
        if len(p) <= cap * 0.8 and len(p) <= cap_line * lines:
            best = p
    return best


_chart_counter = [0]


def content_for(layout_key, style, icons):
    slots = {s["idx"]: s for s in LAYOUTS["styles"][style][layout_key]["slots"]}
    spec = {"layout": layout_key, "style": style, "speaker_notes": "Числа иллюстративные."}
    n = _chart_counter[0]
    for role, idxs in ROLES[layout_key]["roles"].items():
        if role == "title":
            spec[role] = "Итоги года" if layout_key == "section" else "Развитие сети 5G"
        elif role == "chart":
            spec["chart"] = {"type": ["column", "line", "pie", "bar", "stacked_column", "doughnut"][n % 6],
                             "categories": ["2025", "2026", "2027"],
                             "series": [{"name": "БС 5G", "values": [8, 12.5, 18]},
                                        {"name": "LTE", "values": [30, 31, 31.5]}][:1 if n % 6 in (2, 5) else 2],
                             "number_format": "0.0"}
            _chart_counter[0] += 1
            n += 1
        elif role == "picture":
            spec["picture"] = {"icon": icons[len(icons) // 2 + n % 3], "color": "white" if style == "green" else "green"}
        elif isinstance(idxs, int):
            spec[role] = fit_text(slots[idxs])
        elif isinstance(idxs[0], list):
            spec[role] = [[fit_text(slots[i]) for i in row] for row in idxs]
        else:
            spec[role] = [fit_text(slots[i]) for i in idxs]
    if layout_key == "blank":
        spec["title"] = "Свободный слайд"
        spec["elements"] = [
            {"kind": "table", "x": 0.6, "y": 1.3, "w": 5, "h": 1.5, "rows": [["Регион", "БС"], ["Север", "1200"]]},
            {"kind": "chart", "x": 0.6, "y": 3.0, "w": 4, "h": 2.2,
             "chart": {"type": "column", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}},
            {"kind": "text", "x": 5.8, "y": 1.3, "w": 3.5, "h": 0.6, "text": "Подпись", "pt": 14, "color": "purple"},
            {"kind": "icon", "x": 6, "y": 3, "size": 0.8, "name": "antenna", "color": "berry"}]
    return spec


def full_deck():
    icons = sorted(B.icon_names())
    slides = []
    for style in ("light", "green", "dark"):
        for key in ROLES:
            slides.append(content_for(key, style, icons))
    return {"title": "5G и развитие сети", "date": "14.10.2026", "slides": slides}


class BuildTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.path = os.path.join(cls.tmp, "all.pptx")
        cls.deck = full_deck()
        B.validate(cls.deck)
        B.build(cls.deck, cls.path)
        cls.prs = Presentation(cls.path)

    def test_every_layout_every_style(self):
        self.assertEqual(len(self.prs.slides), 3 * len(ROLES))
        used = {s.slide_layout.name for s in self.prs.slides}
        self.assertGreaterEqual(len(used), len(ROLES) - 1)

    def test_charts_tables_pictures(self):
        kinds = set()
        for s in self.prs.slides:
            for sh in s.shapes:
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    kinds.add(str(sh.chart.chart_type))
                if getattr(sh, "has_table", False) and sh.has_table:
                    kinds.add("table")
                if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
                    kinds.add("picture")
        self.assertGreaterEqual(len([k for k in kinds if k not in ("table", "picture")]), 3, kinds)
        self.assertIn("table", kinds)
        self.assertIn("picture", kinds)

    def test_date_and_slide_number(self):
        for n, s in enumerate(self.prs.slides, 1):
            layout_has_dt = any("DATE" in str(p.placeholder_format.type) for p in s.slide_layout.placeholders)
            if not layout_has_dt:
                continue
            dts = [sh for sh in s.placeholders if "DATE" in str(sh.placeholder_format.type)]
            self.assertEqual(len(dts), 1, "slide %d" % n)
            self.assertIn("5G и развитие сети", dts[0].text_frame.text)
            self.assertIn("14.10.2026", dts[0].text_frame.text)
            flds = s._element.findall(".//" + qn("a:fld"))
            self.assertTrue(any(f.get("type") == "slidenum" for f in flds), "slide %d no slidenum field" % n)

    def test_notes(self):
        self.assertEqual(self.prs.slides[2].notes_slide.notes_text_frame.text, "Числа иллюстративные.")

    def test_lint_clean(self):
        errs = [i for i in LD.lint(self.path) if i["level"] == "ERROR"]
        self.assertEqual(errs, [], "\n".join("%s %s %s" % (i["slide"], i["shape"], i["message"]) for i in errs))

    def test_lint_catches_problems(self):
        deck = {"title": "t", "date": "1", "slides": [
            {"layout": "cover", "title": "Обложка"},
            {"layout": "two_pills", "title": "Мегафон - \"лучший\"", "items": ["Очень длинный текст " * 20]}]}
        # bypass typesetting by writing raw text after build
        path = os.path.join(self.tmp, "bad.pptx")
        B.build(deck, path)
        prs = Presentation(path)
        for sh in prs.slides[1].placeholders:
            if sh.placeholder_format.idx == 0:
                sh.text_frame.paragraphs[0].runs[0].text = 'Мегафон - "лучший"'
                sh.text_frame.paragraphs[0].runs[0].font.name = "Calibri"
        prs.save(path)
        msgs = " | ".join(i["message"] for i in LD.lint(path))
        for needle in ("font 'Calibri'", "typesetting", "overflow"):
            self.assertIn(needle, msgs)


class ValidationTests(unittest.TestCase):
    def check(self, slide, needle):
        with self.assertRaises(B.ValidationError) as cm:
            B.validate({"slides": [slide]})
        self.assertIn(needle, str(cm.exception))

    def test_bad_layout(self):
        self.check({"layout": "nope"}, "unknown layout")

    def test_bad_role_lists_allowed(self):
        self.check({"layout": "kpi_2x2", "foo": "x"}, "Allowed roles")

    def test_too_many_items(self):
        self.check({"layout": "two_pills", "items": ["a", "b", "c"]}, "only 2 slots")

    def test_missing_chart(self):
        self.check({"layout": "chart_wide", "title": "x"}, "missing chart data")

    def test_bad_icon_suggestions(self):
        self.check({"layout": "icon_bg_list", "picture": {"icon": "antena"}}, "antenna")

    def test_cli_exit_code(self):
        with tempfile.TemporaryDirectory() as t:
            p = os.path.join(t, "d.json")
            with open(p, "w") as f:
                json.dump({"slides": [{"layout": "nope"}]}, f)
            r = subprocess.run([sys.executable, os.path.join(SCRIPTS, "build_deck.py"), p, "-o",
                                os.path.join(t, "o.pptx")], capture_output=True)
            self.assertEqual(r.returncode, 2)


if __name__ == "__main__":
    unittest.main()
