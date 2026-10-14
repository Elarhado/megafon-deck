#!/usr/bin/env python3
"""Render one contact sheet per style: every layout with each placeholder labelled #idx.
Usage: render_slotmaps.py assets_dir"""
import json, subprocess, sys, tempfile, glob, os
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

A = Path(sys.argv[1]); out = A / "slotmaps"; out.mkdir(exist_ok=True)
cat = json.loads((A / "layouts.json").read_text())
font = ImageFont.truetype(next(iter(glob.glob("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")), None) or "Arial.ttf", 22)
for style, lays in cat["styles"].items():
    prs = Presentation(A / "template.pptx")
    keys = sorted(lays, key=lambda k: lays[k]["index"])
    for k in keys:
        L = prs.slide_masters[lays[k]["master"]].slide_layouts[lays[k]["index"]]
        s = prs.slides.add_slide(L)
        for ph in s.placeholders:
            idx = ph.placeholder_format.idx
            if ph.has_text_frame:
                ph.text_frame.text = f"#{idx}"
            else:
                try:
                    ph.insert_picture  # picture placeholder: leave prompt
                except Exception:
                    pass
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "m.pptx"; prs.save(p)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(p)], capture_output=True, timeout=300)
        subprocess.run(["pdftoppm", "-r", "60", "-png", str(Path(td) / "m.pdf"), str(Path(td) / "pg")], check=True)
        pngs = sorted(glob.glob(str(Path(td) / "pg-*.png")))
        w, h = Image.open(pngs[0]).size; cols = 4; rows = (len(pngs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * (w + 8), rows * (h + 34)), "white"); d = ImageDraw.Draw(sheet)
        for i, f in enumerate(pngs):
            x = (i % cols) * (w + 8); y = (i // cols) * (h + 34)
            d.text((x + 4, y + 4), keys[i], fill="black", font=font)
            sheet.paste(Image.open(f), (x, y + 32))
        sheet.save(out / f"{style}.png", optimize=True)
        print(style, sheet.size)
