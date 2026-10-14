#!/usr/bin/env python3
"""One-off: turn the source MegaFon template into the plugin's clean master template.

- removes sample slides (and the think-cell object on them)
- recolours theme/layouts to brandbook 2026 values (B359FF -> AA5AFF)
- writes assets/layouts.json: catalogue of layouts per style with placeholder slots
Usage: prepare_template.py SOURCE.pptx OUT_DIR
"""
import json, re, shutil, sys, zipfile, os
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

RECOLOR = {"B359FF": "AA5AFF"}            # template -> brandbook 2026
STYLES = {0: "dark", 1: "green", 2: "light"}

# canonical keys by normalised layout name; M2 uses generic names for 3 layouts
KEYS = [
    (r"^Обложка|^13_Заголовок", "cover"),
    (r"^Шмутц|^17_Заголовок", "section"),
    (r"плашк\w*, два подзаголовка\. ?1", "rows_arrows"),
    (r"плашк\w*, два подзаголовка\. ?2", "rows_bars"),
    (r"большой иконкой на фоне\. 2", "icon_bg_two_columns"),
    (r"большой иконкой на фоне\.$", "icon_bg_list"),
    (r"тремя подзаголовками|^46_Заголовок", "table_3col"),
    (r"^Таймлайн", "timeline"),
    (r"^Плашка с текстом", "plate_list"),
    (r"^Две плашки с выводом", "two_plates_conclusion"),
    (r"^Цитаты", "quote"),
    (r"рамочк\w* с текстом", "three_frames"),
    (r"рамочк\w* с выводом", "three_steps_conclusion"),
    (r"^Три скруглённые рамки\.$", "three_pills"),
    (r"^Четыре скруглённые рамки", "four_arrows_conclusion"),
    (r"^Две скруглённые рамки", "two_pills"),
    (r"^Три скруглённые рамки и изображение", "three_pills_image"),
    (r"^Рамка и изображение", "frame_image"),
    (r"^Рамка с текстом", "kpi_2x2"),
    (r"^Диаграмма\.1", "chart_two_items"),
    (r"^Диаграмма\.2", "chart_five_items"),
    (r"^Диаграмма\.3", "chart_wide"),
    (r"^Фон\.?1?$|^Фон\.$", "blank"),
    (r"^Фон\.2", "blank_2"),
    (r"^Фон\.3", "blank_3"),
]
PH_TYPES = {1: "title", 2: "body", 3: "ctr_title", 4: "subtitle", 7: "object", 8: "chart",
            13: "slide_number", 16: "date", 15: "footer", 18: "picture", 12: "table"}


def key_for(name):
    for rx, k in KEYS:
        if re.search(rx, name.strip()):
            return k
    return None


def inch(v):
    return round(Emu(v).inches, 2) if v is not None else None


def font_size(ph):
    # first explicit size in placeholder lstStyle / runs, in pt
    xml = ph._element.xml
    m = re.search(r'sz="(\d+)"', xml)
    return int(m.group(1)) / 100 if m else None


def main(src, out):
    out = Path(out); out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(src)
    # drop sample slides
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
    catalogue = {"slide_size_in": [inch(prs.slide_width), inch(prs.slide_height)], "styles": {}}
    for mi, master in enumerate(prs.slide_masters):
        style = STYLES.get(mi, f"m{mi}")
        lay = {}
        for li, layout in enumerate(master.slide_layouts):
            k = key_for(layout.name)
            if not k:
                print("unmapped layout:", mi, li, layout.name); continue
            slots = []
            for ph in layout.placeholders:
                pf = ph.placeholder_format
                t = PH_TYPES.get(int(pf.type), str(pf.type))
                prompt = ph.text_frame.text.strip()[:60] if ph.has_text_frame else ""
                slots.append({"idx": pf.idx, "type": t, "name": ph.name,
                              "x": inch(ph.left), "y": inch(ph.top), "w": inch(ph.width), "h": inch(ph.height),
                              "pt": font_size(ph), "prompt": prompt})
            slots.sort(key=lambda s: (s["y"] if s["y"] is not None else 0, s["x"] or 0))
            lay[k] = {"layout_name": layout.name, "master": mi, "index": li, "slots": slots}
        catalogue["styles"][style] = lay
    tmp = out / "_tmp.pptx"; prs.save(tmp)
    # recolour at XML level
    dst = out / "template.pptx"
    with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.endswith(".xml"):
                s = data.decode("utf8")
                for a, b in RECOLOR.items():
                    s = re.sub(a, b, s, flags=re.I)
                data = s.encode("utf8")
            zo.writestr(it, data)
    tmp.unlink()
    (out / "layouts.json").write_text(json.dumps(catalogue, ensure_ascii=False, indent=1), encoding="utf8")
    print("styles:", {k: len(v) for k, v in catalogue["styles"].items()}, "->", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
